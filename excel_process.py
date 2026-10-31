"""Перенос часов из нарядов цехов в «Расшифровку ЗП».

Источники (ищутся в текущей папке по имени файла):
  * «наряд» — наряды цехов, по листу на рабочего, месяц берётся из имени файла;
  * «план»  — план производства, из него берутся названия заказов;
  * «зп»    — расшифровка ЗП, в неё записываются часы.
"""

import datetime
import json
import os
import re
from collections import defaultdict
from dataclasses import dataclass, field

from openpyxl import load_workbook
from python_calamine import CalamineWorkbook

curr_dir = os.getcwd()
info_dir = "/Информация полученная в процессе работы/"

MONTHS = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)

# Листы нарядов, которые не относятся к рабочим
SERVICE_SHEETS = {"часы", "пример бланка па", "бланк4"}

PATTERN_SOURCE = "наряд"
PATTERN_PLAN = "план"
PATTERN_DEST = "зп"

WORKSHOP_ORDER = "цех"
# Заказы оснастки (23 счёт) идут в ЗП первыми
TOOLING_PREFIX = "00 07"

# Разметка листа наряда
OUTFIT_MONTH_MARKER = "наряд за месяц"
OUTFIT_ORDER_COL = 1
OUTFIT_HOURS_COL = 3

# Разметка листа плана производства
PLAN_ORDER_HEADER = "№ заказа"
PLAN_NAME_HEADER = "Наименование детали, узла, продукции, вида ремонта"

# Разметка листа «Расшифровки ЗП» (номера строк и столбцов как в Excel, с 1)
ZP_ORDER_NAME_ROW = 4
ZP_ORDER_CODE_ROW = 5
ZP_FIRST_WORKER_ROW = 9
ZP_NAME_COL = 2
ZP_TOTAL_COL = 5
ZP_FIRST_ORDER_COL = 8
ZP_ORDER_COL_STEP = 2  # каждый заказ — пара столбцов «часы | доля»
ZP_WORKSHOP_COL = 272

ORDER_CODE_RE = re.compile(r"\d{7,8}", re.ASCII)


@dataclass
class MonthData:
    # ФИО (normalize_name) -> код заказа -> часы
    workers: dict = field(
        default_factory=lambda: defaultdict(lambda: defaultdict(datetime.timedelta))
    )
    # то же, но для строк, код которых не удалось распознать
    not_included: dict = field(
        default_factory=lambda: defaultdict(lambda: defaultdict(datetime.timedelta))
    )
    # код заказа -> наименование из плана производства
    orders: dict = field(default_factory=dict)
    # ФИО (normalize_name) -> ФИО как оно записано в наряде, для отчётов
    names: dict = field(default_factory=dict)


def normalize_name(name: str) -> str:
    """Ключ для сопоставления ФИО: «Иванов И. И.», «Иванов И.И», «иванов ИИ» -> «ивановии»."""
    return re.sub(r"[\s.]+", "", name).lower()


def normalize_order_code(raw: str) -> str | None:
    """Приводит код заказа к виду «53 01 1 243».

    Пробелы внутри кода не важны: «53 01 1 243», «53 01 1243», «5301 1243», «53011243»
    дают одно и то же. Семизначный код считается кодом с потерянным ведущим нулём.
    """
    digits = re.sub(r"\s+", "", raw)
    if not ORDER_CODE_RE.fullmatch(digits):
        return None
    digits = digits.zfill(8)
    return f"{digits[0:2]} {digits[2:4]} {digits[4]} {digits[5:8]}"


def parse_hours(value) -> datetime.timedelta | None:
    if isinstance(value, datetime.timedelta):
        return value
    if isinstance(value, datetime.time):
        return datetime.timedelta(
            hours=value.hour, minutes=value.minute, seconds=value.second
        )
    if isinstance(value, str) and value.strip():
        try:
            return parse_hours(datetime.time.fromisoformat(value.strip()))
        except ValueError:
            return None
    return None


def to_hours(duration: datetime.timedelta) -> float:
    return round(duration.total_seconds() / 3600, 2)


def cell_text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        # Код заказа, набранный без пробелов, Excel хранит числом
        return str(int(value))
    return value.strip() if isinstance(value, str) else ""


def serialize_datetime(obj):
    if isinstance(obj, datetime.timedelta):
        return str(obj)


def save_json(file_name: str, data):
    with open(curr_dir + info_dir + file_name, "w", encoding="utf8") as f:
        f.write(
            json.dumps(data, default=serialize_datetime, ensure_ascii=False, indent=4)
        )


def find_month(text: str) -> str | None:
    text = text.lower()
    return next((month for month in MONTHS if month in text), None)


def get_name_files() -> list[str]:
    return [
        f
        for f in os.listdir(curr_dir)
        if os.path.isfile(os.path.join(curr_dir, f))
        and f.lower().endswith((".xlsx", ".xls"))
        # «~$...» — временные файлы, которые Excel создаёт рядом с открытой книгой
        and not f.startswith("~$")
    ]


def filter_name_files(files: list[str], pattern: str) -> list[str]:
    return [f for f in files if re.search(pattern, f.lower())]


def find_single_file(files: list[str], pattern: str) -> str:
    found = filter_name_files(files, pattern)
    if len(found) != 1:
        raise Exception(
            f"Должен быть ровно один файл с «{pattern}» в имени, найдено: {found}"
        )
    return found[0]


def read_executor_name(rows: list, sh_name: str, filename: str) -> str:
    """ФИО из ячейки справа от «Исполнитель:» в шапке наряда."""
    surname = sh_name.split(" ")[0].lower()
    for row in rows:
        cells = [cell_text(value) for value in row if cell_text(value)]
        for ind, text in enumerate(cells[:-1]):
            if "исполнитель" in text.lower():
                name = cells[ind + 1]
                if surname not in name.lower():
                    raise Exception(
                        f"Имя листа {sh_name} в файле {filename} не совпадает с {name}"
                    )
                return name
    raise Exception(
        f"На листе {sh_name} в файле {filename} не найдена ячейка «Исполнитель»"
    )


def read_month_outfit(rows: list, sh_name: str, filename: str):
    """Строки (код заказа, часы) из итогового «Наряда за месяц» внизу листа."""
    start = next(
        (
            ind
            for ind, row in enumerate(rows)
            if any(
                cell_text(value).lower().startswith(OUTFIT_MONTH_MARKER)
                for value in row
            )
        ),
        None,
    )
    if start is None:
        raise Exception(
            f"На листе {sh_name} в файле {filename} не найден «Наряд за месяц»"
        )
    for row in rows[start + 1 :]:
        if len(row) <= OUTFIT_HOURS_COL:
            continue
        order = cell_text(row[OUTFIT_ORDER_COL])
        if order.lower() == "итого":
            break
        if order and order != PLAN_ORDER_HEADER:
            yield order, row[OUTFIT_HOURS_COL]


def read_outfits(files: list[str]) -> dict[str, MonthData]:
    months: dict[str, MonthData] = defaultdict(MonthData)
    for file in files:
        month = find_month(file)
        if month is None:
            raise Exception(f"В имени файла {file} не указан месяц")
        data = months[month]
        wb = CalamineWorkbook.from_path(file)
        for sh_name in wb.sheet_names:
            if sh_name.lower() in SERVICE_SHEETS:
                continue
            rows = wb.get_sheet_by_name(sh_name).to_python(skip_empty_area=False)
            # Один рабочий может быть в нарядах нескольких цехов — часы суммируются
            name = read_executor_name(rows[:7], sh_name, file)
            key = normalize_name(name)
            data.names.setdefault(key, name)
            for raw_order, raw_hours in read_month_outfit(rows, sh_name, file):
                hours = parse_hours(raw_hours)
                if hours is None:
                    print(
                        f"Пропущена строка без времени: {file}, лист {sh_name}, заказ {raw_order}"
                    )
                    continue
                order = normalize_order_code(raw_order)
                if order:
                    data.workers[key][order] += hours
                    data.orders.setdefault(order, "")
                elif WORKSHOP_ORDER in raw_order.lower():
                    data.workers[key][WORKSHOP_ORDER] += hours
                else:
                    data.not_included[key][raw_order] += hours
    return months


def read_order_names(file: str, months: dict[str, MonthData]):
    wb = CalamineWorkbook.from_path(file)
    sheets = {sh_name.lower().strip(): sh_name for sh_name in wb.sheet_names}
    for month, data in months.items():
        if month not in sheets:
            print(
                f"В плане производства нет листа «{month}», названия заказов не заполнены"
            )
            continue
        rows = wb.get_sheet_by_name(sheets[month]).to_python(skip_empty_area=False)
        header_ind = next(
            ind for ind, row in enumerate(rows) if PLAN_ORDER_HEADER in row
        )
        header = rows[header_ind]
        order_col, name_col = (
            header.index(PLAN_ORDER_HEADER),
            header.index(PLAN_NAME_HEADER),
        )
        for row in rows[header_ind + 1 :]:
            order = normalize_order_code(cell_text(row[order_col]))
            if order in data.orders:
                data.orders[order] = row[name_col]


def order_columns(orders: dict[str, str]) -> dict[str, int]:
    codes = sorted(orders, key=lambda code: (not code.startswith(TOOLING_PREFIX), code))
    columns = {
        code: ZP_FIRST_ORDER_COL + ind * ZP_ORDER_COL_STEP
        for ind, code in enumerate(codes)
    }
    if columns and max(columns.values()) >= ZP_WORKSHOP_COL:
        raise Exception(
            f"Заказов ({len(columns)}) больше, чем столбцов в ЗП до столбца «цех»"
        )
    columns[WORKSHOP_ORDER] = ZP_WORKSHOP_COL
    return columns


def write_salary_sheet(ws, data: MonthData) -> dict:
    """Записывает часы в лист ЗП и возвращает отчёт о несовпавших ФИО."""
    columns = order_columns(data.orders)
    hours_columns = range(ZP_FIRST_ORDER_COL, ZP_WORKSHOP_COL + 1, ZP_ORDER_COL_STEP)

    # Шапка: стираем заказы прошлых запусков и пишем текущие
    for col in hours_columns[:-1]:
        ws.cell(row=ZP_ORDER_NAME_ROW, column=col).value = None
        ws.cell(row=ZP_ORDER_CODE_ROW, column=col).value = None
    for order, col in columns.items():
        if order != WORKSHOP_ORDER:
            ws.cell(row=ZP_ORDER_NAME_ROW, column=col, value=data.orders[order])
            ws.cell(row=ZP_ORDER_CODE_ROW, column=col, value=order)

    zp_names_not_found = []
    written = set()
    row = ZP_FIRST_WORKER_ROW
    while (
        isinstance(name := ws.cell(row=row, column=ZP_NAME_COL).value, str)
        and name.lower() != "итого"
    ):
        key = normalize_name(name)
        if key in data.workers:
            written.add(key)
            for col in hours_columns:
                ws.cell(row=row, column=col).value = None
            worker_orders = data.workers[key]
            for order, hours in worker_orders.items():
                ws.cell(row=row, column=columns[order], value=to_hours(hours))
            ws.cell(
                row=row,
                column=ZP_TOTAL_COL,
                value=to_hours(sum(worker_orders.values(), datetime.timedelta())),
            )
        else:
            zp_names_not_found.append(name)
        row += 1

    outfit_names_not_found = [
        data.names[key] for key in data.workers if key not in written
    ]
    if not zp_names_not_found and not outfit_names_not_found:
        return {"Всё окей по мнению программы": None}
    return {
        "Имена взяты из ЗП, которым не нашли пару": zp_names_not_found,
        "Имена взяты из нарядов, которым не нашли пару": outfit_names_not_found,
        "Пояснение": [
            "Регистр, пробелы и точки в ФИО не важны, остальное должно совпадать буква в букву, ",
            "данные берутся из наряда справа-сверху в ячейке исполнитель.",
            "",
            "Просто сверь имена взятые из ЗП с именами взятыми из нарядов",
            "",
            "Самый редкий случай, когда кто-то вместо русской С добавил английскую C",
            "Разницу не заметит никто, кроме программы",
        ],
    }


def write_salary(file: str, months: dict[str, MonthData]) -> dict:
    wb = load_workbook(file)
    sheets = {sh_name.lower().strip(): sh_name for sh_name in wb.sheetnames}
    report = {}
    for month, data in months.items():
        if month not in sheets:
            print(f"В файле ЗП нет листа «{month}», месяц пропущен")
            continue
        report[month] = write_salary_sheet(wb[sheets[month]], data)
    wb.save(file)
    return report


def sorted_by_key(d: dict) -> dict:
    return dict(sorted(d.items()))


def run():
    os.makedirs(curr_dir + info_dir, exist_ok=True)
    files = get_name_files()

    print("Загрузка данных о рабочих началась")
    source_files = filter_name_files(files, PATTERN_SOURCE)
    if not source_files:
        raise Exception("Не найден файл с именем наряд")
    months = read_outfits(source_files)
    print("Загрузка данных о рабочих окончилась")

    print("Загрузка номеров заказов из плана производства началась")
    read_order_names(find_single_file(files, PATTERN_PLAN), months)
    print("Загрузка номеров заказов из плана производства окончилась")

    save_json(
        "Все заказы в общем и заказы каждого рабочего.json",
        {
            month: {
                **{
                    data.names[key]: sorted_by_key(orders)
                    for key, orders in sorted_by_key(data.workers).items()
                },
                "ords": sorted_by_key(data.orders),
            }
            for month, data in months.items()
        },
    )
    save_json(
        "Не вошедшие в ЗП заказы.json",
        {
            month: {
                data.names[key]: dict(orders)
                for key, orders in data.not_included.items()
            }
            for month, data in months.items()
        },
    )

    print("Редактирование файла ЗП началось")
    report = write_salary(find_single_file(files, PATTERN_DEST), months)
    save_json("ФИО которых нет в ЗП.json", report)
    print("Редактирование файла ЗП окончилось")
